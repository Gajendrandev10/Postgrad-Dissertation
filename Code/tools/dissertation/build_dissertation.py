"""Build with the bundled document Python runtime; preserve the MMU template."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import re
import zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parents[3]/"Documentation/02_Dissertation/support"
ROOT=HERE.parents[1]
TEMPLATE=ROOT/'04_Templates/Word/mmu_dissertation_template.docx'
FINAL=ROOT/'02_Dissertation/uav_cybersecurity_dissertation_draft.docx'
EXPECTED='cc30ae24038fa48e0c9385173010714e92832239429b5c58902aac63fcdd6dad'
assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()==EXPECTED
TEXT=(HERE/'dissertation_draft.md').read_text()
E=json.loads((HERE/'evidence_analysis.json').read_text())
PAGES=json.loads((HERE/'page_map.json').read_text()) if (HERE/'page_map.json').exists() else {}
D=Document(TEMPLATE)
# Reuse the template's institutional cover logo drawing and package relationship.
logo=next(deepcopy(p._p) for p in D.paragraphs if p._p.xpath('.//w:drawing'))
body=D._element.body
for c in list(body):
    if c.tag!=qn('w:sectPr'): body.remove(c)
# All sample document/header/footer text is an editable template slot.
for part in list(D.part.package.parts):
    if str(part.partname).startswith(('/word/header','/word/footer')):
        el=part.element
        for c in list(el): el.remove(c)
        el.append(OxmlElement('w:p'))
section=D.sections[0]
for tag in ['headerReference','footerReference','titlePg','pgNumType']:
    for e in list(section._sectPr.findall(qn('w:'+tag))): section._sectPr.remove(e)

def style(name,size=12,bold=False,space=6,line=1.5):
    s=D.styles[name] if name in D.styles else D.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Calibri';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k),'Calibri')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+k),None)
    pf=s.paragraph_format;pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.space_before=Pt(0);pf.space_after=Pt(space);pf.line_spacing=line
    pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(0);pf.widow_control=True
    ppr=s.element.get_or_add_pPr()
    for k in ['numPr','pBdr','tabs']:
        for x in list(ppr.findall(qn('w:'+k))): ppr.remove(x)
    return s

normal=style('Normal')
for name in ['Body Text','Body Text 2','Body Text 3']: style(name)
for n,size in [(1,18),(2,14),(3,12)]:
    s=style('Heading '+str(n),size,True);s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=(n==1)
    ppr=s.element.get_or_add_pPr();o=ppr.find(qn('w:outlineLvl'))
    if o is None:o=OxmlElement('w:outlineLvl');ppr.append(o)
    o.set(qn('w:val'),str(n-1))
for name in ['Front heading','Index heading']:
    s=style(name,18,True);s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=True
cap=style('Caption',11,False);cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.keep_together=True
style('Table text',10.5,False,4,1.5)
style('TOC 1',12,False,6,1.5);style('TOC 2',12,False,3,1.5).paragraph_format.left_indent=Cm(.5)
style('Reference',12,False,6,1.5).paragraph_format.left_indent=Cm(.95)
D.styles['Reference'].paragraph_format.first_line_indent=Cm(-.95)

def geometry(s):
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.54)
    s.header_distance=Cm(1.25);s.footer_distance=Cm(1.25)
    s.gutter=Cm(0)
    for el in s._sectPr.findall(qn('w:cols')):el.set(qn('w:num'),'1')

def field(p,code,cached='1'):
    r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    r=p.add_run();inst=OxmlElement('w:instrText');inst.set(qn('xml:space'),'preserve');inst.text=' '+code+' ';r._r.append(inst)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    p.add_run(str(cached));r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)

BOOKS={};next_id=200
def bookmark(p,key):
    global next_id
    next_id+=1;name='uav_'+key;BOOKS[key]=name
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(next_id));b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(next_id))
    p._p.insert(1 if p._p.pPr is not None else 0,b);p._p.append(e)

def paragraph(text='',sty=None):
    p=D.add_paragraph(text,sty);return p

def heading(text,level=1,key=None,front=False):
    p=paragraph(text,'Front heading' if front else 'Heading '+str(level))
    if key:bookmark(p,key)
    return p

def footer(s,fmt,start):
    s.footer.is_linked_to_previous=False
    p=s.footer.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(p,'PAGE',str(start))
    pg=s._sectPr.find(qn('w:pgNumType'))
    if pg is None:pg=OxmlElement('w:pgNumType');s._sectPr.append(pg)
    pg.set(qn('w:fmt'),fmt);pg.set(qn('w:start'),str(start))

def add_index(title,code,entries):
    h=paragraph(title,'Index heading')
    first=True;last=None
    for label,key,level in entries:
        p=paragraph('', 'TOC 2' if level==2 else 'TOC 1')
        p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        if first:
            for kind in ['begin','code','separate']:
                r=p.add_run()
                if kind=='code': el=OxmlElement('w:instrText');el.set(qn('xml:space'),'preserve');el.text=' '+code+' '
                else:el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
                r._r.append(el)
            first=False
        p.add_run(label+'\t');field(p,'PAGEREF uav_'+key+' \\h',PAGES.get(key,'1'));last=p
    if last:
        r=last.add_run();el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),'end');r._r.append(el)
    return h

geometry(section)
title='UAV communication anomaly detection and SOC integration with a post-quantum cryptographic demonstration'
p=paragraph();p.paragraph_format.space_before=Pt(42)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(title);r.bold=True;r.font.size=Pt(18)
paragraph()
for t in ['A dissertation for the degree of Master of Science','Manchester Metropolitan University','[Exact programme title]']:
    p=paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
paragraph()
p=paragraph();p._p.addprevious(logo)
logo_pr=logo.get_or_add_pPr()
for e in list(logo_pr): logo_pr.remove(e)
jc=OxmlElement('w:jc');jc.set(qn('w:val'),'center');logo_pr.append(jc)
logo.xpath('.//wp:docPr')[0].set('descr','Manchester Metropolitan University logo retained from the supplied template')
for e in logo.xpath('.//wp:extent'):
    e.set('cx',str(int(Cm(3.2))));e.set('cy',str(int(Cm(3.2))))
for e in logo.xpath('.//a:xfrm/a:ext'):
    e.set('cx',str(int(Cm(3.2))));e.set('cy',str(int(Cm(3.2))))
for t in ['September 2026','[Full name]','Student ID: [Student ID]','Supervisor: [Supervisor name]']:
    p=paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER

pre=D.add_section(WD_SECTION_START.NEW_PAGE);geometry(pre);footer(pre,'lowerRoman',1)
abstract=('UAV security analysis requires traceable observations and a clear distinction between a simulation result and an operational security claim. '
'This dissertation develops and evaluates a local prototype linking OMNeT++ records, Isolation Forest and Random Forest predictions, structured security events and a separate ML-KEM-768 with AES-256-GCM proof of concept. '
'A preserved five-UAV simulation produces 247 send records: 220 normal and 27 abnormal. A stratified split reserves 75 records for testing. Random Forest classifies all 75 correctly; Isolation Forest detects all eight test attacks and produces five false positives, giving 61.54% precision and 100% recall. '
'The union policy generates 13 events and adds no true detection beyond Random Forest on this split. The generator places normal and abnormal latency and loss in non-overlapping ranges, and a generator-informed threshold also separates the test records perfectly. '
'The cryptographic benchmark completes 200 verified local round trips, with mean key-generation, encapsulation and decapsulation times of 0.01681, 0.01729 and 0.02056 milliseconds. '
'Output hashes and source identifiers support inspection of the preserved experiment, although the editable notebook has subsequently diverged. Wazuh-oriented export is implemented; live Wazuh delivery and per-event encrypted transport are not demonstrated. '
'The findings support a reproducible controlled proof of concept. Independent scenarios, benign communication disturbances, authenticated transport and target-device measurements are required before making deployment claims.')
h=heading('Abstract',key='abstract',front=True);h.paragraph_format.page_break_before=False
paragraph(abstract)
heading('Declaration',key='declaration',front=True)
for t in ['[Author to complete and verify the originality declaration using the applicable university wording.]',
          'EthOS approval number: [Genuine approval number]',
          '[Author to record assistance and sources as required by university policy. The template statement that the work is unaided has not been asserted.]',
          'Signed: __________________________', 'Date: __________________________']:
    paragraph(t)
heading('Abbreviations',key='abbreviations',front=True)
for a,b in [('AES','Advanced Encryption Standard'),('CSV','Comma-separated values'),('EDI','Equality, diversity and inclusion'),('FANET','Flying ad hoc network'),('GCM','Galois/Counter Mode'),('HKDF','HMAC-based extract-and-expand key derivation function'),('IF','Isolation Forest'),('KEM','Key-encapsulation mechanism'),('ML-KEM','Module-lattice-based key-encapsulation mechanism'),('NDJSON','Newline-delimited JSON'),('PQC','Post-quantum cryptography'),('RF','Random Forest'),('ROC','Receiver operating characteristic'),('SOC','Security operations centre'),('UAV','Unmanned aerial vehicle')]:
    p=paragraph(a+'\t'+b);p.paragraph_format.tab_stops.add_tab_stop(Cm(2.5))

entries=[('Abstract','abstract',1),('Declaration','declaration',1),('Abbreviations','abbreviations',1)]
blocks=TEXT.split('\n\n')
key_by_title={}
for b in blocks:
    if b.startswith('# '):
        t=b[2:].strip();key='ch'+t.split(':')[0].split()[-1];entries.append((t,key,1));key_by_title[t]=key
    elif b.startswith('## '):
        t=b[3:].strip();key='s'+t.split()[0].replace('.','_');entries.append((t,key,2));key_by_title[t]=key
entries += [('References','references',1),('Appendix A: Terms of Reference','appendix_a',1),('Appendix B: Ethics approval','appendix_b',1),('Appendix C: Digital artefact repository','appendix_c',1),('Appendix D: Evidence and reproduction guide','appendix_d',1)]
add_index('Contents','TOC \\o "1-2" \\t "Front heading,1" \\h \\z',entries)
figs=[];tabs=[]
for b in blocks:
    m=re.fullmatch(r'\[\[(FIGURE|TABLE):(\w+)\|(.+)\]\]',b.strip())
    if m:
        kind,key,captext=m.groups();target=figs if kind=='FIGURE' else tabs
        number=len(target)+1;target.append((f'{kind.title()} {number}: {captext}',('fig_' if kind=='FIGURE' else 'tab_')+key,1))
add_index('List of tables','TOC \\c "Table" \\h \\z',tabs)
add_index('List of figures','TOC \\c "Figure" \\h \\z',figs)

main=D.add_section(WD_SECTION_START.NEW_PAGE);geometry(main);footer(main,'decimal',1)

TABLES={
'dataset':(['Population','Normal','Attack','Total'],[
['Raw events','467','27','494'],['Retained send records','220','27','247'],['Training records','153','19','172'],['Test records','67','8','75']], [7.5,2.8,2.8,2.8]),
'evidence':(['Evidence file','Evaluation role'],[
['run_manifest.json','Input/output hashes and run identity'],['executed_notebook.ipynb','Executed experiment code and cell output'],['omnet_uav_ai_dataset.csv','247 selected send records and features'],['split_membership.csv','Train/test allocation by source identifier'],['uav_test_predictions.csv','75 held-out model decisions and scores'],['soc_alerts.csv / uav_soc_events.json','13 alert records and NDJSON export'],['pqc_benchmark_results.csv','200 local timing and success records'],['environment.json','Recorded platform and Python package versions']], [7.4,8.5]),
'metrics':(['Model','Accuracy\n(%)','Precision\n(%)','Recall\n(%)','F1\n(%)','FPR\n(%)'],[
['Isolation Forest','93.33','61.54','100.00','76.19','7.46'],['Random Forest','100.00','100.00','100.00','100.00','0.00']], [4.4,2.3,2.3,2.3,2.3,2.3]),
'pqc':(['Operation','Mean','Median','p95','Maximum'],[
['Key generation']+[f'{E["pqc"]["keygen_ms"][k]:.5f}' for k in ['mean','median','p95','max']],
['Encapsulation']+[f'{E["pqc"]["encapsulation_ms"][k]:.5f}' for k in ['mean','median','p95','max']],
['Decapsulation']+[f'{E["pqc"]["decapsulation_ms"][k]:.5f}' for k in ['mean','median','p95','max']],
['AES encryption']+[f'{E["pqc"]["aes_encrypt_ms"][k]:.5f}' for k in ['mean','median','p95','max']],
['AES decryption']+[f'{E["pqc"]["aes_decrypt_ms"][k]:.5f}' for k in ['mean','median','p95','max']]], [5.1,2.7,2.7,2.7,2.7]),
'sizes':(['Object','Bytes','Meaning'],[
['ML-KEM public key','1,184','Recorded in every trial'],['KEM ciphertext','1,088','Recorded in every trial'],['Shared secret','32','Local secret material; not transmitted'],['Plaintext','299','Representative serialised payload'],['Encrypted message','315','Ciphertext plus authentication tag'],['Nonce','12','Separate value generated by the code'],['Message plus nonce','327','Calculated; excludes protocol framing']], [5.1,2,8.8])}

def blank(keep=False):
    p=paragraph();p.paragraph_format.keep_with_next=keep
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.5
    p.add_run().font.size=Pt(12)
    return p

def caption(kind,n,key,txt):
    p=paragraph('', 'Caption');p.add_run(kind+' ');field(p,'SEQ '+kind+' \\* ARABIC',n);p.add_run(': '+txt)
    bookmark(p,('fig_' if kind=='Figure' else 'tab_')+key)
    return p

counts={'FIGURE':0,'TABLE':0}
for b in blocks:
    b=b.strip()
    if not b:continue
    if b.startswith('# '):
        t=b[2:];p=heading(t,key=key_by_title[t])
        if t.startswith('Chapter 1:'):p.paragraph_format.page_break_before=False
    elif b.startswith('## '):
        t=b[3:];heading(t,2,key_by_title[t])
    elif b.startswith('[[FIGURE:') or b.startswith('[[TABLE:'):
        m=re.fullmatch(r'\[\[(FIGURE|TABLE):(\w+)\|(.+)\]\]',b);kind,key,txt=m.groups();counts[kind]+=1
        blank(True)
        if kind=='FIGURE':
            p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(HERE/'figures'/f'{key}.png'),width=Cm(15.7))
            drawing=p._p.xpath('.//wp:docPr')[0];drawing.set('descr',txt)
        else:
            headers,rows,widths=TABLES[key]
            # Scale to the usable A4 width.
            widths=[w*15.7/sum(widths) for w in widths]
            table=D.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            for c,w in zip(table.columns,widths):c.width=Cm(w)
            for c,h,w in zip(table.rows[0].cells,headers,widths):c.text=h;c.width=Cm(w)
            for row in rows:
                for c,t,w in zip(table.add_row().cells,row,widths):c.text=t;c.width=Cm(w)
            pr=table._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tblCellMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for i,row in enumerate(table.rows):
                trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if i==0:trpr.append(OxmlElement('w:tblHeader'))
                for j,c in enumerate(row.cells):
                    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if i==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF0');c._tc.get_or_add_tcPr().append(shade)
                    for p in c.paragraphs:
                        p.style='Table text';p.paragraph_format.keep_with_next=True
                        if key in ['dataset','metrics','pqc'] and j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:r.bold=(i==0)
        caption(kind.title(),counts[kind],key,txt);blank()
    else:paragraph(b)

heading('References',key='references')
REFS=[('National Institute of Standards and Technology, Module-Lattice-Based Key-Encapsulation Mechanism Standard, FIPS 203, 2024, doi: 10.6028/NIST.FIPS.203.', 'https://doi.org/10.6028/NIST.FIPS.203'), ('I. Bekmezci, O. K. Sahingoz, and S. Temel, "Flying ad-hoc networks (FANETs): A survey," Ad Hoc Networks, vol. 11, no. 3, pp. 1254-1270, 2013, doi: 10.1016/j.adhoc.2012.12.004.', 'https://doi.org/10.1016/j.adhoc.2012.12.004'), ('OMNeT++ Community, "OMNeT++ simulation manual."', 'https://doc.omnetpp.org/omnetpp/manual/'), ('S. Mishra, B. Bhargava, Z. Liu, and S. Islam, "UAV-CAS: A calibrated digital-twin dataset for intrusion detection in UAV swarm networks," arXiv:2606.17845, ver. 1, 2026, preprint, doi: 10.48550/arXiv.2606.17845.', 'https://doi.org/10.48550/arXiv.2606.17845'), ('F. T. Liu, K. M. Ting, and Z.-H. Zhou, "Isolation forest," in Proc. 8th IEEE Int. Conf. Data Mining, 2008, pp. 413-422.', 'https://research.monash.edu/en/publications/isolation-forest/'), ('L. Breiman, "Random forests," Machine Learning, vol. 45, pp. 5-32, 2001, doi: 10.1023/A:1010933404324.', 'https://doi.org/10.1023/A:1010933404324'), ('R. Sommer and V. Paxson, "Outside the closed world: On using machine learning for network intrusion detection," in Proc. IEEE Symp. Security and Privacy, 2010, pp. 305-316.', 'https://oaklandsok.github.io/papers/sommer2010.pdf'), ('F. Pendlebury, F. Pierazzi, R. Jordaney, J. Kinder, and L. Cavallaro, "TESSERACT: Eliminating experimental bias in malware classification across space and time," in Proc. 28th USENIX Security Symp., 2019, pp. 729-746.', 'https://www.usenix.org/conference/usenixsecurity19/presentation/pendlebury'), ('scikit-learn developers, "Common pitfalls and recommended practices," ver. 1.9 documentation.', 'https://scikit-learn.org/stable/common_pitfalls.html'), ('Canadian Institute for Cybersecurity, "Intrusion detection evaluation dataset (CIC-IDS2017)," University of New Brunswick.', 'https://www.unb.ca/cic/datasets/ids-2017.html'), ('M. Dworkin, Recommendation for Block Cipher Modes of Operation: Galois/Counter Mode (GCM) and GMAC, NIST SP 800-38D, 2007, doi: 10.6028/NIST.SP.800-38D.', 'https://doi.org/10.6028/NIST.SP.800-38D'), ('G. Alagic et al., Recommendations for Key-Encapsulation Mechanisms, NIST SP 800-227, 2025, doi: 10.6028/NIST.SP.800-227.', 'https://doi.org/10.6028/NIST.SP.800-227'), ('H. Krawczyk and P. Eronen, "HMAC-based Extract-and-Expand Key Derivation Function (HKDF)," RFC 5869, 2010.', 'https://www.rfc-editor.org/rfc/rfc5869'), ('Open Quantum Safe, "liboqs."', 'https://openquantumsafe.org/liboqs/'), ('Wazuh, "Data analysis," User manual.', 'https://documentation.wazuh.com/current/user-manual/ruleset/index.html'), ('Wazuh, "Testing decoders and rules," User manual.', 'https://documentation.wazuh.com/current/user-manual/ruleset/testing.html'), ('scikit-learn developers, "IsolationForest," ver. 1.9 documentation.', 'https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.IsolationForest.html'), ('NIST/SEMATECH, "7.2.4.1 Confidence intervals," e-Handbook of Statistical Methods.', 'https://www.itl.nist.gov/div898/handbook/prc/section2/prc241.htm'), ('Information Commissioner\'s Office, "Unmanned Aerial Systems (UAS) / Drones."', 'https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/cctv-and-video-surveillance/guidance-on-video-surveillance-including-cctv/additional-considerations-for-technologies-other-than-cctv/unmanned-aerial-systems-uas-drones/'), ('Crown Prosecution Service, "Computer Misuse Act," Prosecution guidance.', 'https://www.cps.gov.uk/prosecution-guidance/computer-misuse-act'), ('Association for Computing Machinery, ACM Code of Ethics and Professional Conduct, 2018.', 'https://www.acm.org/binaries/content/assets/about/acm-code-of-ethics-booklet.pdf')]
for ref_number,(txt,url) in enumerate(REFS,1):
    p=paragraph('['+str(ref_number)+']\t'+txt+' Accessed: Sep. 6, 2026. [Online]. Available: ', 'Reference')
    p.paragraph_format.tab_stops.add_tab_stop(Cm(.95))
    # Standard external hyperlink; readable URL wraps naturally in Word.
    link=OxmlElement('w:hyperlink');rid=D.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);link.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'000000');rp.append(co);r.append(rp);t=OxmlElement('w:t');t.text=url;r.append(t);link.append(r);p._p.append(link)
    p.add_run('.')

heading('Appendix A: Terms of Reference',key='appendix_a')
paragraph('[Insert the approved Terms of Reference, including its original aims, objectives, scope and approval details.]')
heading('Appendix B: Ethics approval',key='appendix_b')
paragraph('[Insert the genuine MMU EthOS approval email. Enter its actual approval number in the declaration and retain the identifying approval details here.]')
heading('Appendix C: Digital artefact repository',key='appendix_c')
paragraph('Repository URL: [Examiner-accessible repository link]')
paragraph('Release or commit identifier: [Identifier of the assessed artefact]')
paragraph('Access verification: [Date and method used to confirm examiner access]')
paragraph('[Confirm that the assessed code and results remain accessible through at least the end of the assessment period in mid-November, as required by the assignment brief.]')
heading('Appendix D: Evidence and reproduction guide',key='appendix_d')
paragraph('The empirical reference is the preserved directory 06_Results/runs/local_only_pqc_20260906. It contains the executed notebook, model inputs and predictions, event exports, timing measurements and environment record. Relative paths in this appendix start at the project root.')
paragraph('Raw input: 01_OMNeT_Simulation/UAV_SOC_Project/simulations/uav_traffic.csv. SHA-256:')
p=paragraph('819966a8705aa359936ba02e69c304f65c61e805b32c9007d728032e61cacbc2');p.paragraph_format.keep_together=True
paragraph('The reference directory contains run_manifest.json with input and output digests. Its executed_notebook.ipynb is the implementation source used in this dissertation where the mutable notebook differs.')
paragraph('Verify output hashes, 247 prepared records, 172 training identifiers, 75 test identifiers and 13 events. Join predictions to raw rows using source_record_id. Confirm disjoint partitions and prediction membership in the test set.')
paragraph('The supporting script 02_Dissertation/support/analyse_saved_run.py recomputes the descriptive quantities used in Chapter 5 without retraining or changing the saved run. It records output and input hash checks, confusion counts, event membership, timing summaries and illustrative Wilson intervals in evidence_analysis.json.')
paragraph('The editable notebook fails the manifest input-hash check and includes Colab-related differences. Reconcile it with the preserved source before rerunning. Use a fresh output directory and retain a separate identifier, environment and manifest for the new run.')
paragraph('The local PQC installation context is documented in 03_PQC_Security/local_installation.json and local_macos_setup.md. The native liboqs source commit is 8c7f6592e097736c414b08eea12817b29399734f; the wrapper commit is 0f9e50bea240792b802a6e56f5e4d33a9ec0a6b0. Reproduction must use a compatible native architecture and verify algorithm availability before benchmarking.')
paragraph('The legacy_run directories contain incompatible 795-record results. They must not supply values for the reference run. The supplied templates and examples are presentation references only.')

# Settings for genuine Word automatic caption numbering and index refresh.
settings=D.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
for x in settings.findall(qn('w:evenAndOddHeaders')):settings.remove(x)
captions=settings.find(qn('w:captions'))
if captions is None:captions=OxmlElement('w:captions');settings.append(captions)
for kind in ['Figure','Table']:
    el=OxmlElement('w:caption');el.set(qn('w:name'),kind);el.set(qn('w:numFmt'),'decimal');captions.append(el)
for s in D.sections:
    geometry(s);s.different_first_page_header_footer=False
D.core_properties.title=title;D.core_properties.subject='Controlled UAV simulation, AI/SOC evaluation and local PQC demonstration'
D.core_properties.author='';D.core_properties.last_modified_by='';D.core_properties.comments=''
buffer=BytesIO();D.save(buffer)
# Preserve unrelated template package parts byte-for-byte.
source=zipfile.ZipFile(TEMPLATE);built=zipfile.ZipFile(BytesIO(buffer.getvalue()))
preserved=[]
with zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as z:
    for n in built.namelist():
        preserve=(n.startswith(('word/media/','word/theme/','customXml/')) or n in ['word/fontTable.xml','word/numbering.xml','word/footnotes.xml','word/endnotes.xml']) and n in source.namelist()
        z.writestr(n,source.read(n) if preserve else built.read(n))
        if preserve:preserved.append(n)
assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()==EXPECTED
alltext=' '.join(D.element.xpath('//w:t/text()'))
report={'template_sha256':EXPECTED,'preserved_parts':preserved,'chapters':7,'main_source_word_count':len(TEXT.split()),'visible_word_count_approx':len(alltext.split()),'references':len(REFS),'figures':counts['FIGURE'],'tables':counts['TABLE'],'bookmarks':BOOKS,'output':str(FINAL),'field_refresh':'updateFields=true; refresh in Word before final submission','main_headings':key_by_title}
(HERE/'build_report.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k not in ['bookmarks','main_headings','preserved_parts']},indent=2))
