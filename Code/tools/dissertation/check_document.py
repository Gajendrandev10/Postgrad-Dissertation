"""Structural, evidence and rendered-text checks for the Word draft."""
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

HERE=Path(__file__).resolve().parents[3]/"Documentation/02_Dissertation/support"
ROOT=HERE.parents[1]
DOC=ROOT/'02_Dissertation/uav_cybersecurity_dissertation_draft.docx'
REF=ROOT/'04_Templates/Word/mmu_dissertation_template.docx'
d=Document(DOC)
b=json.loads((HERE/'build_report.json').read_text())
assert hashlib.sha256(REF.read_bytes()).hexdigest()==b['template_sha256']
with zipfile.ZipFile(DOC) as z,zipfile.ZipFile(REF) as s:
    assert all(z.read(n)==s.read(n) for n in b['preserved_parts'])
for sec in d.sections:
    assert abs(sec.page_width.cm-21)<.01 and abs(sec.page_height.cm-29.7)<.01
    assert all(abs(x.cm-2.54)<.01 for x in [sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin])
for name,size in [('Normal',12),('Heading 1',18),('Heading 2',14),('Heading 3',12),('Caption',11)]:
    style=d.styles[name]
    assert style.font.name=='Calibri' and style.font.size.pt==size
    assert str(style.font.color.rgb)=='000000'
assert d.styles['Normal'].paragraph_format.line_spacing==1.5
assert d.styles['Normal'].paragraph_format.space_before.pt==0
assert d.styles['Normal'].paragraph_format.space_after.pt==6
assert d.settings.element.find(qn('w:updateFields')).get(qn('w:val'))=='true'
fields=d.element.xpath('//w:instrText/text()')
assert sum('SEQ Figure' in x for x in fields)==4
assert sum('SEQ Table' in x for x in fields)==5
assert sum(x.strip().startswith('TOC') for x in fields)==3
text=' '.join(d.element.xpath('//w:t/text()'))
assert not re.search('[—–]',text)
assert 'THE TITLE OF THE PROJECT' not in text and 'Your EthOS Number' not in text
main=[];active=False
for p in d.paragraphs:
    if p.text=='Chapter 1: Introduction':active=True
    if p.text=='References':active=False
    if active:main.append(p.text)
actual_main_count=len(' '.join(main).split())
assert 10000<=actual_main_count<=15000
with pdfplumber.open(sys.argv[1]) as pdf:
    fonts=sorted({c['fontname'] for p in pdf.pages for c in p.chars})
    assert all('Calibri' in x for x in fonts)
    texts=[p.extract_text() or '' for p in pdf.pages]
    assert texts[1].startswith('Abstract') and 'deployment claims.' in texts[1]
    assert texts[2].startswith('Declaration')
    assert all('Error! Reference' not in t and 'Error: Reference' not in t for t in texts)
    overflows=[]
    for n,p in enumerate(pdf.pages,1):
        chars=[c for c in p.chars if c.get('text','').strip()]
        if any(c['x0']<70 or c['x1']>p.width-70 for c in chars):overflows.append(n)
    assert not overflows,overflows
    pages=len(pdf.pages)
# IEEE numbering must match first appearance and every source must be cited.
body_text=' '.join(main)
citation_numbers=[int(n) for n in re.findall(r'\[(\d+)\]',body_text)]
assert list(dict.fromkeys(citation_numbers))==list(range(1,22))
reference_paragraphs=[p for p in d.paragraphs if p.style.name=='Reference']
assert [int(re.match(r'\[(\d+)\]',p.text).group(1)) for p in reference_paragraphs]==list(range(1,22))
assert len(citation_numbers)==22
for entry in json.loads((HERE/'ieee_citation_map.json').read_text()):
    assert entry['old_citation'] not in body_text
    assert entry['new_citation'] in body_text
q={
    'docx_sha256':hashlib.sha256(DOC.read_bytes()).hexdigest(),
    'template_unchanged':True,'preserve_only_parts_unchanged':True,
    'page_count':pages,'fonts':fonts,'main_body_word_count_excluding_tables':actual_main_count,
    'approx_complete_visible_word_count':b['visible_word_count_approx'],
    'margins_and_styles_pass':True,'abstract_one_page':True,
    'figures':4,'tables':5,'references':21,'citation_style':'IEEE','in_text_citation_count':22,'citation_numbering_pass':True,'caption_seq_fields':9,'index_toc_fields':3,
    'horizontal_text_bounds_pass':True,'rendered_pdf':str(Path(sys.argv[1]).resolve()),
    'visual_review':'Pending individual page inspection',
    'submission_ready':False,
    'remaining_author_inputs':['Personal details','Confirmed title and programme','Approved Terms of Reference','Genuine EthOS email and number','Examiner-accessible repository link','Verified declaration'],
    'evidence_limitation':'All saved output hashes match; the editable notebook input hash differs substantively.'
}
(HERE/'qa_report.json').write_text(json.dumps(q,indent=2)+'\n')
print(json.dumps(q,indent=2))
